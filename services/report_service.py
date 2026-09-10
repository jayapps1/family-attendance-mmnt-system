from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4
from xml.sax.saxutils import escape
from decimal import Decimal
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
from models import ContributionPayment
from repositories.contribution_repository import ContributionRepository
from services.base import Service, snapshot
from services.family_service import FamilyService
from services.relationship_service import RelationshipService
from services.meeting_service import MeetingService
from services.attendance_service import AttendanceService
from services.contribution_service import ContributionService
from services.settings_service import SettingsService
from utils.validators import ValidationError


class ReportService(Service):
    REPORTS = (
        "Family register", "Individual member", "Family branch", "Genealogy", "Meetings",
        "Attendance", "Member attendance history", "Contribution summary", "Fully paid members",
        "Partially paid members", "Unpaid members", "Outstanding balances",
        "Payment transaction history", "Individual contribution history",
    )

    def __init__(self, sessions, actor_id, output_root):
        super().__init__(sessions, actor_id)
        self.output_root = Path(output_root)

    def data(self, name, *, member_id=None, period_id=None, affiliation_type=None):
        if name not in self.REPORTS:
            raise ValidationError("Unknown report")
        family = FamilyService(self.sessions, self.actor_id)
        members = {row["id"]: row for row in family.list()}
        from utils.family_labels import affiliation_label
        names = {i: f'{r["family_number"]} - {r["first_name"]} {r["last_name"]}' for i, r in members.items()}
        if name in {"Family register", "Individual member", "Family branch"}:
            rows = list(members.values())
            if name == "Individual member":
                if member_id is None:
                    raise ValidationError("Select a member")
                rows = [members[member_id]]
            elif name == "Family branch":
                if member_id is None:
                    raise ValidationError("Select the branch's founding member")
                tree = RelationshipService(self.sessions, self.actor_id).tree(member_id)
                ids = {member_id} | {r["id"] for r in tree["descendants"]}
                rows = [r for r in rows if r["id"] in ids]
            from models import FamilyBranch
            from repositories.base import Repository
            links = RelationshipService(self.sessions, self.actor_id)
            with links.transaction() as session:
                graph = links.graph(session)
                branches = Repository(session, FamilyBranch).list()
                enriched = []
                for r in rows:
                    if affiliation_type is not None and r['affiliation_type'] != affiliation_type:
                        continue
                    ancestors = {r['id']} | set(graph.depths(r['id'], ancestors=True))
                    enriched.append(dict(r, affiliation=affiliation_label(r['affiliation_type']),
                        birth_position=graph.members[r['id']].get('birth_position', ''),
                        branch='; '.join(b.name for b in branches if b.founding_member_id in ancestors)))
                rows = enriched
            columns = ("family_number", "first_name", "last_name", "affiliation", "birth_position", "branch", "sex", "age", "phone_number", "current_residence", "living_status")
        elif name == "Genealogy":
            rows = RelationshipService(self.sessions, self.actor_id).list()
            rows = [dict(parent=names.get(r["parent_id"], ""), child=names.get(r["child_id"], ""),
                         relationship=r["relationship_type"]) for r in rows]
            columns = ("parent", "child", "relationship")
        elif name == "Meetings":
            rows = MeetingService(self.sessions, self.actor_id).list()
            columns = ("meeting_number", "title", "meeting_date", "venue", "status")
        elif name in {"Attendance", "Member attendance history"}:
            attendance = AttendanceService(self.sessions, self.actor_id)
            if name == "Member attendance history" and member_id is None:
                raise ValidationError("Select a member")
            rows = attendance.member_history(member_id) if member_id else attendance.list()
            meetings = {r["id"]: r["title"] for r in MeetingService(self.sessions, self.actor_id).list()}
            rows = [dict(r, member=names.get(r["family_member_id"], ""), meeting=meetings.get(r["meeting_id"], "")) for r in rows]
            columns = ("member", "meeting", "status", "reason", "arrival_time")
        elif name == "Contribution summary":
            contributions = ContributionService(self.sessions, self.actor_id)
            periods = contributions.periods()
            rows = [dict(period=p["title"], **contributions.daily_summary(p["id"], affiliation_type)) for p in periods
                    if period_id is None or p["id"] == period_id]
            columns = ("period", "expected_total", "collected", "outstanding", "PAID", "PARTIALLY_PAID", "UNPAID", "eligible_contributors", "UNDER_23", "DECEASED", "DOB_UNKNOWN", "LIVING_UNKNOWN")
        elif name == "Payment transaction history":
            with self.transaction() as session:
                repo = ContributionRepository(session)
                obligations = {r.id: r for r in repo.list()}
                payments = repo.payments.list()
                rows = [dict(snapshot(p), member=names.get(obligations[p.member_contribution_id].family_member_id, ""))
                        for p in payments if (period_id is None or obligations[p.member_contribution_id].contribution_period_id == period_id)
                        and (affiliation_type is None or members[obligations[p.member_contribution_id].family_member_id]['affiliation_type'] == affiliation_type)]
            columns = ("receipt_number", "member", "payment_date", "amount_paid", "payment_method", "is_reversed", "reversal_reason")
        else:
            contributions = ContributionService(self.sessions, self.actor_id)
            if name == "Individual contribution history" and member_id is None:
                raise ValidationError("Select a member")
            status = {"Fully paid members": "PAID", "Partially paid members": "PARTIALLY_PAID",
                      "Unpaid members": "UNPAID"}.get(name)
            if name == "Individual contribution history":
                rows = contributions.obligations(period_id, member_id, status, affiliation_type)
            else:
                rows = [row for period in contributions.periods()
                        if period_id is None or period["id"] == period_id
                        for row in contributions.register_rows(period["id"])
                        if (status is None or row["status"] == status)
                        and (affiliation_type is None or row["affiliation_type"] == affiliation_type)
                        and (row["historical"] or row["eligible"])]
            periods = {p["id"]: p["title"] for p in contributions.periods()}
            rows = [dict(r, member=names.get(r["family_member_id"], ""), period=periods.get(r["contribution_period_id"], "")) for r in rows]
            if name == "Outstanding balances":
                rows = [r for r in rows if r["outstanding"] > 0]
            rows = [dict(r, affiliation=affiliation_label(r["affiliation_type"])) for r in rows]
            columns = ("member", "affiliation", "period", "amount_due", "total_paid", "outstanding", "status")
        return columns, rows

    def export(self, name, format, **filters):
        columns, rows = self.data(name, **filters)
        settings = {r["setting_key"]: r["setting_value"] for r in SettingsService(self.sessions, self.actor_id).list()}
        title = settings.get("report_header") or settings.get("family_name") or "Family Management"
        footer = settings.get("report_footer", "")
        folder = self.output_root / ("pdf" if format == "PDF" else "excel")
        if format not in {"PDF", "Excel"}:
            raise ValidationError("Select PDF or Excel")
        folder.mkdir(parents=True, exist_ok=True)
        filename = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S_") + uuid4().hex[:8]
        path = folder / (filename + (".pdf" if format == "PDF" else ".xlsx"))
        headers = [key.replace("_", " ").title() for key in columns]
        if format == "Excel":
            book = Workbook()
            sheet = book.active
            sheet.title = "Report"
            sheet.append([title])
            sheet.append([name])
            sheet.append(headers)
            for row in rows:
                sheet.append([row.get(key) for key in columns])
            # User-controlled strings must remain text, never spreadsheet formulas.
            for cells in sheet:
                for cell in cells:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
                    if isinstance(cell.value, Decimal):
                        cell.number_format = '#,##0.00'
            sheet.freeze_panes = "A4"
            sheet.auto_filter.ref = f"A3:{sheet.cell(max(3, sheet.max_row), len(columns)).coordinate}"
            from openpyxl.utils import get_column_letter
            for i in range(1, len(columns) + 1):
                sheet.column_dimensions[get_column_letter(i)].width = 24
            if footer:
                sheet.append([footer])
                sheet.cell(sheet.max_row, 1).data_type = "s"
            book.save(path)
        else:
            styles = getSampleStyleSheet()
            body = styles["BodyText"]
            body.fontSize, body.leading = 8, 10
            paragraph = lambda value: Paragraph(escape(str(value if value is not None else "")), body)
            table = LongTable([[paragraph(h) for h in headers]] +
                              [[paragraph(row.get(key)) for key in columns] for row in rows],
                              repeatRows=1, colWidths=[770 / len(columns)] * len(columns))
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dce8ef")),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f7fa")]),
            ]))
            SimpleDocTemplate(str(path), pagesize=landscape(A4), leftMargin=30, rightMargin=30).build([
                Paragraph(escape(title), styles["Title"]), Paragraph(escape(name), styles["Heading2"]),
                Spacer(1, 12), table, Spacer(1, 12), paragraph(footer)])
        with self.transaction() as session:
            from services.audit_service import append_audit
            append_audit(session, self.actor_id, "EXPORT_REPORT", "report", description=name)
        return str(path.resolve())
